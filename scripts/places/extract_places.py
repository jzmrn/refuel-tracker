"""Extract German municipalities (Gemeinden) with coordinates into a CSV file.

Source dataset: "Gemeindeverzeichnis - Quartalsausgabe (Auszug GV 2Q)" published by
the Statistisches Bundesamt (Destatis):

    https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/Administrativ/Archiv/GVAuszugQ/AuszugGV2QAktuell.html

The workbook contains a hierarchical listing (Satzart 10 = Bundesland,
40 = Kreis, 50 = Gemeindeverband, 60 = Gemeinde). Only the Satzart 60 rows carry
the geographic centre coordinates, so those rows are exported while the state and
district names are resolved from their parent rows.

Usage:
    # Use a previously downloaded workbook
    python extract_places.py --input ~/Downloads/AuszugGV2QAktuell.xlsx

    # Download the current release from Destatis
    python extract_places.py --download

    # Custom output location
    python extract_places.py --download --output /tmp/places.csv
"""

from __future__ import annotations

import argparse
import csv
import logging
import sys
import tempfile
import urllib.request
from dataclasses import asdict, dataclass, fields
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger("extract_places")

DOWNLOAD_URL = (
    "https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/"
    "Gemeindeverzeichnis/Administrativ/Archiv/GVAuszugQ/"
    "AuszugGV2QAktuell.xlsx?__blob=publicationFile"
)

DEFAULT_OUTPUT = (
    Path(__file__).resolve().parents[2] / "backend" / "app" / "resources" / "places.csv"
)

# Sheet holding the actual data is named "Onlineprodukt_Gemeinden<date>"
DATA_SHEET_PREFIX = "Onlineprodukt_Gemeinden"

# Column indices (0-based) of the data sheet
COL_SATZART = 0
COL_LAND = 2
COL_RB = 3
COL_KREIS = 4
COL_VB = 5
COL_GEM = 6
COL_NAME = 7
COL_AREA = 8
COL_POPULATION = 9
COL_POSTAL_CODE = 13
COL_LONGITUDE = 14
COL_LATITUDE = 15

# Satzart values of the hierarchy levels we care about
SATZART_STATE = "10"
SATZART_DISTRICT = "40"
SATZART_MUNICIPALITY = "60"

# First row containing data (rows 1-6 are titles and multi-level headers)
FIRST_DATA_ROW = 7


@dataclass(frozen=True)
class Place:
    """A single German municipality with its geographic centre."""

    ars: str
    name: str
    designation: str
    postal_code: str
    district: str
    state: str
    lat: float
    lng: float
    population: int


def download_workbook(target: Path) -> Path:
    """Download the current GV quarterly workbook from Destatis."""
    logger.info("Downloading dataset from %s", DOWNLOAD_URL)
    request = urllib.request.Request(
        DOWNLOAD_URL,
        headers={
            # Destatis rejects requests without a common browser user agent
            "User-Agent": (
                "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            ),
            "Accept": "*/*",
        },
    )
    with urllib.request.urlopen(request) as response, target.open("wb") as handle:
        handle.write(response.read())
    logger.info("Saved workbook to %s (%d bytes)", target, target.stat().st_size)
    return target


def _clean(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _split_name(raw_name: str) -> tuple[str, str]:
    """Split "Flensburg, Stadt" into ("Flensburg", "Stadt").

    Destatis appends the administrative designation (Stadt, St, M, Flecken, ...)
    after the last comma. Municipality names themselves never contain a comma.
    """
    if "," in raw_name:
        name, designation = raw_name.rsplit(",", 1)
        return name.strip(), designation.strip()
    return raw_name, ""


def _parse_coordinate(value: object) -> float | None:
    """Parse a coordinate cell, which uses a German decimal comma."""
    text = _clean(value).replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _parse_int(value: object) -> int:
    text = _clean(value).replace(".", "").replace(" ", "")
    try:
        return int(float(text))
    except ValueError:
        return 0


def _find_data_sheet(workbook) -> str:
    for sheet_name in workbook.sheetnames:
        if sheet_name.startswith(DATA_SHEET_PREFIX):
            return sheet_name
    raise ValueError(
        f"No sheet starting with '{DATA_SHEET_PREFIX}' found. "
        f"Available sheets: {workbook.sheetnames}"
    )


def extract_places(workbook_path: Path) -> list[Place]:
    """Read the workbook and return all municipalities that have coordinates."""
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_name = _find_data_sheet(workbook)
        logger.info("Reading sheet '%s' from %s", sheet_name, workbook_path)
        sheet = workbook[sheet_name]

        places: list[Place] = []
        current_state = ""
        current_district = ""
        skipped = 0

        for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
            satzart = _clean(row[COL_SATZART])

            if satzart == SATZART_STATE:
                current_state = _split_name(_clean(row[COL_NAME]))[0]
                current_district = ""
                continue

            if satzart == SATZART_DISTRICT:
                current_district = _split_name(_clean(row[COL_NAME]))[0]
                continue

            if satzart != SATZART_MUNICIPALITY:
                continue

            lat = _parse_coordinate(row[COL_LATITUDE])
            lng = _parse_coordinate(row[COL_LONGITUDE])
            raw_name = _clean(row[COL_NAME])

            if lat is None or lng is None or not raw_name:
                skipped += 1
                logger.debug("Skipping row without coordinates: %s", raw_name)
                continue

            name, designation = _split_name(raw_name)
            ars = "".join(
                _clean(row[column])
                for column in (COL_LAND, COL_RB, COL_KREIS, COL_VB, COL_GEM)
            )

            places.append(
                Place(
                    ars=ars,
                    name=name,
                    designation=designation,
                    postal_code=_clean(row[COL_POSTAL_CODE]),
                    district=current_district,
                    state=current_state,
                    lat=lat,
                    lng=lng,
                    population=_parse_int(row[COL_POPULATION]),
                )
            )

        logger.info(
            "Extracted %d places (%d rows without coordinates)", len(places), skipped
        )
        return places
    finally:
        workbook.close()


def write_csv(places: list[Place], output: Path) -> None:
    """Write the extracted places to a UTF-8 CSV file sorted by name."""
    output.parent.mkdir(parents=True, exist_ok=True)
    column_names = [field.name for field in fields(Place)]

    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=column_names)
        writer.writeheader()
        for place in sorted(places, key=lambda p: (p.name.casefold(), p.ars)):
            writer.writerow(asdict(place))

    logger.info("Wrote %d places to %s", len(places), output)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument(
        "--input",
        type=Path,
        help="Path to an already downloaded AuszugGV2QAktuell.xlsx",
    )
    source.add_argument(
        "--download",
        action="store_true",
        help="Download the current workbook from Destatis before extracting",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Target CSV file (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "--verbose",
        action="store_true",
        help="Enable debug logging",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s - %(message)s",
    )

    if args.download:
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = download_workbook(Path(tmp_dir) / "AuszugGV2QAktuell.xlsx")
            places = extract_places(workbook_path)
    else:
        if not args.input.exists():
            logger.error("Input file does not exist: %s", args.input)
            return 1
        places = extract_places(args.input)

    if not places:
        logger.error("No places extracted - the workbook layout may have changed")
        return 1

    write_csv(places, args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
